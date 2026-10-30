"""把 4.2 版 xlsm（9～11 月）轉成資料庫匯入用 xlsx。"""
import sys, datetime as dt, warnings, openpyxl
from openpyxl.styles import Font, PatternFill
warnings.filterwarnings('ignore')
HOURS = ['08','09','10','11','12','13','14','15','16','17','18','19','20','21','22','23','00','01','02','03','04','05','06','07']

def s(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()

def role(n):
    if n == 1: return '分隊長'
    if 2 <= n <= 9: return '小隊長'
    if n in (51, 52): return '實務生'
    if n == 53: return '實習生'
    return '隊員'

people = {}      # 番號 -> {id, 姓名}
roster, archive, calendar_like = [], [], []
first_day = None
for path in sys.argv[2:]:
    wb = openpyxl.load_workbook(path, data_only=True)
    m = wb['月班表']
    year, month = int(s(m.cell(1, 1).value)) + 1911, int(s(m.cell(2, 1).value))
    cols, c = [], 3
    while s(m.cell(1, c).value):
        n = int(s(m.cell(1, c).value)); name = s(m.cell(2, c).value)
        if n not in people:
            people[n] = {'id': f'P{len(people)+1:03d}', 'name': name}
        cols.append((c, n)); c += 1
    for r in range(3, 34):
        d = s(m.cell(r, 1).value)
        if not d.isdigit() or not s(m.cell(r, 2).value): continue
        date = dt.date(year, month, int(d))
        first_day = min(first_day or date, date)
        for c, n in cols:
            roster.append((date, people[n]['id'], s(m.cell(r, c).value)))
    ws = wb['存檔區']; r = 1
    while r <= ws.max_row:
        if s(ws.cell(r, 2).value) == '08':
            date = dt.date(year, month, int(s(ws.cell(r, 1).value))); r += 1
            while r <= ws.max_row and s(ws.cell(r, 2).value) != '08':
                label = s(ws.cell(r, 1).value)
                if label:
                    for h in range(24):
                        v = s(ws.cell(r, 2 + h).value)
                        if v.isdigit() and int(v) in people:
                            archive.append((date, r, label, HOURS[h], people[int(v)]['id'], s(ws.cell(r, 27).value)))
                r += 1
        else:
            r += 1

out = openpyxl.Workbook()
bold = Font(bold=True); head = PatternFill('solid', fgColor='DDEBF7')
def sheet(title, header, rows, first=False):
    ws = out.active if first else out.create_sheet()
    ws.title = title; ws.append(header)
    for cell in ws[1]: cell.font = bold; cell.fill = head
    for row in rows: ws.append(list(row))
    ws.freeze_panes = 'A2'
    for i, h in enumerate(header): ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width = max(10, len(h)*2+4)
    return ws

sheet('人員', ['人員ID','姓名','別名','身分','計入人力','勤務方式','備註'],
      [(p['id'], p['name'], p['name'][1:], role(n), 'Y' if role(n) not in ('實務生','實習生') else 'N', '一般', '') for n, p in sorted(people.items())], first=True)
sheet('番號期間', ['人員ID','番號','起日','迄日'], [(p['id'], n, first_day, None) for n, p in sorted(people.items())])
sheet('異動紀錄', ['人員ID','起日','迄日','狀態','備註'], [])
sheet('月班表', ['日期','人員ID','代號'], roster)
TYPES = [('水源','B','A','F','梯次'),('駕訓','C','B','L',''),('大隊送公文','C','A','P',''),('消防查察','A','A','P',''),
         ('瓦斯路邊攔查','A','A','N',''),('訓練','C','C','L','內容'),('其他勤務','C','A','P','內容'),('宣導','A','A','G','內容'),
         ('開會','C','A','P','內容'),('住警器安裝','C','A','G',''),('第二位值班','D','D','D',''),('第三位91','D','E','E',''),('第三位92','D','F','E','')]
sheet('勤務類別', ['名稱','Z碼','AA碼','AB碼','需輸入','計入服勤時數'], [(n,z,a,b,x,'N' if a=='C' else 'Y') for n,z,a,b,x in TYPES])
sheet('設定', ['項目','值','說明'], [
    ('備勤最低人數', 3, 'C23 兩車三人'), ('駕訓時備勤減少', 2, 'C23'), ('小隊長每日上限時數', 4, 'C22 值班+救護+共同勤務'),
    ('每人休息時數', 2, 'C14'), ('連續服勤上限', 8, 'C03'), ('分隊長不排班', 'Y', 'C21')])
sheet('每日勤務', ['日期','列','勤務名稱','時','人員ID','旗標'], archive)
for ws in out.worksheets:
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, dt.date): cell.number_format = 'yyyy-mm-dd'
out.save(sys.argv[1])
print(len(people), 'people', len(roster), 'roster rows', len(archive), 'archive cells')
