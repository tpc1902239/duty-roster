"""讀取 4.2 版 xlsm：月班表、存檔區（每日排班）、救護訓練。"""
import warnings
import openpyxl

warnings.simplefilter('ignore')
HOURS = ['08', '09', '10', '11', '12', '13', '14', '15', '16', '17', '18', '19',
         '20', '21', '22', '23', '00', '01', '02', '03', '04', '05', '06', '07']
ON_DUTY = {'', 'A1', 'A2', 'A3', 'A4', 'B1', 'B2', 'B3', '離', '宿', '外', '長', 'M'}
FIXED_ROWS = ['值班', '一救', '一救', '二救', '二救', '休息1', '休息2', '休息3', '休息4']


def _s(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    m = wb['月班表']
    people, c = [], 3
    while _s(m.cell(1, c).value):
        people.append(int(_s(m.cell(1, c).value)))
        c += 1
    roster = {}
    for r in range(3, 34):
        d = _s(m.cell(r, 1).value)
        if not d.isdigit() or not _s(m.cell(r, 2).value):
            continue
        roster[int(d)] = {'weekday': _s(m.cell(r, 2).value),
                          'status': {p: _s(m.cell(r, 3 + i).value) for i, p in enumerate(people)}}
    days = {}
    ws = wb['存檔區']
    r = 1
    while r <= ws.max_row:
        if _s(ws.cell(r, 2).value) == '08':
            d = int(_s(ws.cell(r, 1).value))
            rows = []
            r += 1
            while r <= ws.max_row and _s(ws.cell(r, 2).value) != '08':
                label = _s(ws.cell(r, 1).value)
                if label:
                    cells = [int(_s(ws.cell(r, 2 + h).value)) if _s(ws.cell(r, 2 + h).value).isdigit() else None
                             for h in range(24)]
                    rows.append({'label': label, 'cells': cells, 'flag': _s(ws.cell(r, 27).value)})
                r += 1
            days[d] = rows
        else:
            r += 1
    training = {}
    t = wb['救護訓練']
    for r in range(2, t.max_row + 1):
        d, who = _s(t.cell(r, 1).value), _s(t.cell(r, 5).value)
        if d.isdigit() and who.isdigit():
            training[int(d)] = int(who)
    return {'people': people, 'roster': roster, 'days': days, 'training': training}


def night_workers(prev_status):
    """前一天是 B 班或值宿的人（C10：隔天 12 時前不排勤務）。"""
    return {p for p, v in prev_status.items() if v in ('B1', 'B2', 'B3', '宿')}
